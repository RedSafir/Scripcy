"""Comprehensive unit test suite for WordCraft MCP client and DOCX document generation."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from scrpcy.core.mcp_client import (
    CitationItem,
    MockMCPTransport,
    WordCraftClient,
    format_inline_citation_apa,
    format_inline_citation_ieee,
    format_reference_entry_apa,
    format_reference_entry_ieee,
)
from scrpcy.core.models import DocumentMetadata
from scrpcy.core.workspace import WorkspaceManager
from scrpcy.exceptions import (
    DocumentGenerationError,
    MCPConnectionError,
)


class TestCitationFormatting:
    """Tests for APA 7th and IEEE citation text and bibliographic entry formatting."""

    def test_apa_inline_single_author(self) -> None:
        c = CitationItem(authors=["Ashish Vaswani"], year=2017)
        assert format_inline_citation_apa([c]) == "(Vaswani, 2017)"

    def test_apa_inline_two_authors(self) -> None:
        c = CitationItem(authors=["Ashish Vaswani", "Noam Shazeer"], year=2017)
        assert format_inline_citation_apa([c]) == "(Vaswani & Shazeer, 2017)"

    def test_apa_inline_three_plus_authors(self) -> None:
        c = CitationItem(authors=["Ashish Vaswani", "Noam Shazeer", "Niki Parmar"], year=2017)
        assert format_inline_citation_apa([c]) == "(Vaswani et al., 2017)"

    def test_apa_inline_multiple_citations(self) -> None:
        c1 = CitationItem(authors=["Kaiming He", "Xiangyu Zhang"], year=2016)
        c2 = CitationItem(authors=["Ashish Vaswani"], year=2017)
        assert format_inline_citation_apa([c1, c2]) == "(He & Zhang, 2016; Vaswani, 2017)"

    def test_apa_inline_no_author_or_year(self) -> None:
        c = CitationItem()
        assert format_inline_citation_apa([c]) == "(Anonymous, n.d.)"

    def test_ieee_inline_citations(self) -> None:
        assert format_inline_citation_ieee([1]) == "[1]"
        assert format_inline_citation_ieee([2, 1]) == "[1, 2]"
        assert format_inline_citation_ieee([]) == ""

    def test_apa_reference_entry_formatting(self) -> None:
        c = CitationItem(
            title="Attention Is All You Need",
            authors=["Ashish Vaswani", "Noam Shazeer"],
            year=2017,
            doi="10.48550/arXiv.1706.03762",
        )
        entry = format_reference_entry_apa(c)
        assert "Vaswani, A., & Shazeer, N." in entry
        assert "(2017)." in entry
        assert "Attention Is All You Need." in entry
        assert "https://doi.org/10.48550/arXiv.1706.03762" in entry

    def test_ieee_reference_entry_formatting(self) -> None:
        c = CitationItem(
            title="Deep Residual Learning",
            authors=["Kaiming He", "Xiangyu Zhang"],
            year=2016,
            doi="10.1109/CVPR.2016.90",
        )
        entry = format_reference_entry_ieee(c, index=1)
        assert entry.startswith("[1]")
        assert "K. He and X. Zhang" in entry
        assert '"Deep Residual Learning,"' in entry
        assert "2016." in entry
        assert "doi: 10.1109/CVPR.2016.90." in entry


class TestMCPToolDispatching:
    """Tests for MCP JSON-RPC tool dispatch and fallback error boundaries."""

    def test_mcp_tool_dispatch_recorded_calls(self, temp_workspace: WorkspaceManager, tmp_path: Path) -> None:
        mock_transport = MockMCPTransport()
        client = WordCraftClient(
            workspace=temp_workspace,
            transport=mock_transport,
            fallback_to_docx=True,
        )

        assert client.is_using_mcp is True
        doc_path = tmp_path / "mcp_test.docx"

        client.create_document("MCP Research Paper", path=doc_path)
        client.append_heading("Introduction", level=1)
        client.add_paragraph_with_citation(
            "Transformers are self-attention models.",
            citations=[{"authors": ["Ashish Vaswani"], "year": 2017}],
            style="apa",
        )
        client.append_reference_list(
            references=[{"title": "Attention Paper", "authors": ["A. Vaswani"], "year": 2017}],
            style="apa",
        )
        saved = client.save(doc_path)

        assert saved == doc_path
        assert saved.is_file()

        tool_names = [call["params"]["name"] for call in mock_transport.recorded_calls]
        assert "create_document" in tool_names
        assert "append_heading" in tool_names
        assert "add_paragraph_with_citation" in tool_names
        assert "append_reference_list" in tool_names
        assert "save_document" in tool_names

    def test_mcp_failure_clean_error_boundary_and_docx_fallback(
        self,
        temp_workspace: WorkspaceManager,
        tmp_path: Path,
    ) -> None:
        mock_transport = MockMCPTransport(fail_with=MCPConnectionError("Simulated MCP Server Crash"))
        client = WordCraftClient(
            workspace=temp_workspace,
            transport=mock_transport,
            fallback_to_docx=True,
        )

        doc_path = tmp_path / "fallback_paper.docx"

        # Operation should smoothly fall back to local python-docx without crashing
        client.create_document("Fallback Paper", path=doc_path)
        assert client.is_using_mcp is False

        client.append_heading("Background", level=1)
        client.add_paragraph_with_citation(
            "Local fallback ensures high reliability.",
            citations=[{"authors": ["Jane Doe"], "year": 2024}],
            style="apa",
        )
        saved_path = client.save(doc_path)

        assert saved_path.is_file()
        doc = docx.Document(str(saved_path))
        text_content = [p.text for p in doc.paragraphs]
        assert any("Fallback Paper" in t for t in text_content)
        assert any("Background" in t for t in text_content)
        assert any("Local fallback ensures high reliability" in t for t in text_content)

    def test_mcp_failure_raises_when_fallback_disabled(
        self,
        temp_workspace: WorkspaceManager,
        tmp_path: Path,
    ) -> None:
        mock_transport = MockMCPTransport(fail_with=MCPConnectionError("MCP Server Unavailable"))
        client = WordCraftClient(
            workspace=temp_workspace,
            transport=mock_transport,
            fallback_to_docx=False,
        )

        with pytest.raises(MCPConnectionError, match="MCP Server Unavailable"):
            client.create_document("Fail Doc", path=tmp_path / "fail.docx")


class TestDocxDocumentGeneration:
    """Tests verifying headings, paragraphs, and reference formatting in .docx files."""

    def test_generate_apa_docx_document(self, temp_workspace: WorkspaceManager, tmp_path: Path) -> None:
        client = WordCraftClient(workspace=temp_workspace)
        output_file = tmp_path / "apa_paper.docx"

        client.create_document("Empirical Study of LLMs", path=output_file)
        client.append_heading("1. Introduction", level=1)
        client.append_heading("1.1 Background", level=2)

        meta1 = DocumentMetadata(
            title="Attention Is All You Need",
            authors=["Ashish Vaswani", "Noam Shazeer"],
            year=2017,
            doi="10.48550/arXiv.1706.03762",
            cite_key="vaswani2017attention",
        )
        client.add_paragraph_with_citation(
            "The transformer architecture relies entirely on attention.",
            citations=[meta1],
            style="apa",
        )

        client.append_reference_list([meta1], style="apa")
        saved_path = client.save(output_file)

        assert saved_path.is_file()
        doc = docx.Document(str(saved_path))

        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
        assert "Empirical Study of LLMs" in headings
        assert "1. Introduction" in headings
        assert "1.1 Background" in headings
        assert "References" in headings

        paragraphs = [p.text for p in doc.paragraphs if not (p.style.name.startswith("Heading") or p.style.name == "Title")]
        body_para = [p for p in paragraphs if "The transformer architecture" in p]
        assert len(body_para) == 1
        assert "(Vaswani & Shazeer, 2017)" in body_para[0]

        ref_para = [p for p in paragraphs if "Vaswani, A., & Shazeer, N." in p]
        assert len(ref_para) == 1
        assert "https://doi.org/10.48550/arXiv.1706.03762" in ref_para[0]

    def test_generate_ieee_docx_document(self, temp_workspace: WorkspaceManager, tmp_path: Path) -> None:
        client = WordCraftClient(workspace=temp_workspace)
        output_file = tmp_path / "ieee_paper.docx"

        client.create_document("IEEE Benchmark Report", path=output_file)
        client.append_heading("Methods", level=1)

        ref = CitationItem(
            title="Residual Learning",
            authors=["Kaiming He", "Xiangyu Zhang"],
            year=2016,
            doi="10.1109/CVPR.2016.90",
        )

        client.add_paragraph_with_citation(
            "Deep networks benefit from residual connections.",
            citations=[ref],
            style="ieee",
        )
        client.append_reference_list([ref], style="ieee")
        saved_path = client.save(output_file)

        assert saved_path.is_file()
        doc = docx.Document(str(saved_path))
        all_text = [p.text for p in doc.paragraphs]

        assert any("[1]" in t for t in all_text)
        assert any('[1] K. He and X. Zhang, "Residual Learning," 2016. doi: 10.1109/CVPR.2016.90.' in t for t in all_text)

    def test_unrecognized_tool_raises_document_generation_error(self, temp_workspace: WorkspaceManager) -> None:
        client = WordCraftClient(workspace=temp_workspace)
        with pytest.raises(DocumentGenerationError, match="Unrecognized document tool"):
            client.call_tool("invalid_tool_name", {})

    def test_operations_without_active_document_raise_error(self, temp_workspace: WorkspaceManager) -> None:
        client = WordCraftClient(workspace=temp_workspace)
        with pytest.raises(DocumentGenerationError, match="No active document"):
            client.append_heading("Heading before document creation")

    def test_save_without_path_uses_default(self, temp_workspace: WorkspaceManager) -> None:
        client = WordCraftClient(workspace=temp_workspace)
        created_path = client.create_document("Default Path Paper")
        saved_path = client.save()

        assert saved_path == created_path
        assert saved_path.is_file()
