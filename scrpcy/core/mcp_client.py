"""Model Context Protocol (MCP) client and python-docx document generator."""

from __future__ import annotations

from abc import ABC, abstractmethod
import json
from pathlib import Path
import re
import shutil
import subprocess
from typing import Any, Literal

import docx
from pydantic import BaseModel, Field

from scrpcy.core.models import DocumentMetadata
from scrpcy.core.search import SearchResult, sanitize_filename
from scrpcy.core.workspace import WorkspaceManager
from scrpcy.exceptions import (
    DocumentGenerationError,
    MCPConnectionError,
    MCPError,
)

CitationStyle = Literal["apa", "ieee"]


class CitationItem(BaseModel):
    """Normalized citation metadata for academic referencing."""

    title: str = Field(default="", description="Title of the cited work.")
    authors: list[str] = Field(default_factory=list, description="Author names.")
    year: int | None = Field(default=None, description="Year of publication.")
    doi: str | None = Field(default=None, description="Digital Object Identifier.")
    cite_key: str | None = Field(default=None, description="Citation key identifier.")
    url: str | None = Field(default=None, description="Web URL if applicable.")
    institution: str | None = Field(default=None, description="Publishing institution or government agency.")
    access_date: str | None = Field(default=None, description="Date retrieved.")
    source_type: str = Field(default="paper", description="Type of source: 'paper' or 'official_web'.")

    @classmethod
    def from_any(cls, item: Any) -> CitationItem:
        """Convert arbitrary metadata representations into a normalized CitationItem."""
        if isinstance(item, CitationItem):
            return item
        if isinstance(item, (DocumentMetadata, SearchResult)):
            return cls(
                title=item.title,
                authors=item.authors,
                year=item.year,
                doi=item.doi,
                cite_key=item.cite_key,
                url=getattr(item, "url", None),
                institution=getattr(item, "institution", None),
                access_date=getattr(item, "access_date", None),
                source_type=getattr(item, "source_type", "paper"),
            )
        if isinstance(item, dict):
            return cls(
                title=item.get("title", ""),
                authors=item.get("authors", []),
                year=item.get("year"),
                doi=item.get("doi"),
                cite_key=item.get("cite_key"),
                url=item.get("url"),
                institution=item.get("institution"),
                access_date=item.get("access_date"),
                source_type=item.get("source_type", "paper"),
            )
        if isinstance(item, str):
            return cls(title=item, cite_key=item)
        return cls()


def format_author_name_apa(name: str) -> str:
    """Format single author as 'Lastname, Initials' for APA style."""
    parts = name.strip().split()
    if not parts:
        return "Unknown"
    if len(parts) == 1:
        return parts[0]
    last = parts[-1]
    initials = ". ".join(p[0].upper() for p in parts[:-1] if p) + "."
    return f"{last}, {initials}"


def format_author_name_ieee(name: str) -> str:
    """Format single author as 'Initials Lastname' for IEEE style."""
    parts = name.strip().split()
    if not parts:
        return "Unknown"
    if len(parts) == 1:
        return parts[0]
    last = parts[-1]
    initials = ". ".join(p[0].upper() for p in parts[:-1] if p) + "."
    return f"{initials} {last}"


def format_inline_citation_apa(citations: list[CitationItem]) -> str:
    """Format inline citation block following APA 7th edition rules."""
    if not citations:
        return ""

    formatted_entries: list[str] = []
    for item in citations:
        year_str = str(item.year) if item.year else "n.d."
        author_count = len(item.authors)

        if author_count == 0:
            author_str = "Anonymous"
        elif author_count == 1:
            author_str = item.authors[0].strip().split()[-1]
        elif author_count == 2:
            last1 = item.authors[0].strip().split()[-1]
            last2 = item.authors[1].strip().split()[-1]
            author_str = f"{last1} & {last2}"
        else:
            last1 = item.authors[0].strip().split()[-1]
            author_str = f"{last1} et al."

        formatted_entries.append(f"{author_str}, {year_str}")

    return f"({'; '.join(formatted_entries)})"


def format_inline_citation_ieee(indices: list[int]) -> str:
    """Format inline numbered bracket citation following IEEE style."""
    if not indices:
        return ""
    joined = ", ".join(str(i) for i in sorted(indices))
    return f"[{joined}]"


def format_reference_entry_apa(item: CitationItem) -> str:
    """Format single bibliographic reference entry according to APA 7th style."""
    if not item.authors:
        author_str = "Anonymous."
    elif len(item.authors) == 1:
        single = format_author_name_apa(item.authors[0])
        author_str = single if single.endswith(".") else f"{single}."
    elif len(item.authors) == 2:
        a1 = format_author_name_apa(item.authors[0])
        a2 = format_author_name_apa(item.authors[1])
        a2_dot = a2 if a2.endswith(".") else f"{a2}."
        author_str = f"{a1}, & {a2_dot}"
    else:
        formatted_names = [format_author_name_apa(a) for a in item.authors[:-1]]
        last_name = format_author_name_apa(item.authors[-1])
        last_dot = last_name if last_name.endswith(".") else f"{last_name}."
        author_str = f"{', '.join(formatted_names)}, & {last_dot}"

    year_str = f"({item.year})." if item.year else "(n.d.)."
    title_str = f"{item.title.rstrip('.')}." if item.title else "Untitled Document."

    doi_str = ""
    if item.doi:
        clean_doi = item.doi.strip()
        if not clean_doi.startswith("http"):
            clean_doi = f"https://doi.org/{clean_doi}"
        doi_str = f" {clean_doi}"

    if getattr(item, "source_type", None) == "official_web" or (item.url and not item.doi):
        inst_or_auth = item.institution or author_str.rstrip(".")
        retrieved_str = f" Diakses pada {item.access_date}, dari " if getattr(item, "access_date", None) else " "
        return f"{inst_or_auth}. {year_str} {title_str}{retrieved_str}{item.url}"

    return f"{author_str} {year_str} {title_str}{doi_str}"


def format_reference_entry_ieee(item: CitationItem, index: int) -> str:
    """Format single bibliographic reference entry according to IEEE style."""
    if not item.authors:
        author_str = "Anonymous"
    elif len(item.authors) == 1:
        author_str = format_author_name_ieee(item.authors[0])
    elif len(item.authors) == 2:
        author_str = f"{format_author_name_ieee(item.authors[0])} and {format_author_name_ieee(item.authors[1])}"
    else:
        formatted_names = [format_author_name_ieee(a) for a in item.authors[:-1]]
        last_name = format_author_name_ieee(item.authors[-1])
        author_str = f"{', '.join(formatted_names)}, and {last_name}"

    title_str = f'"{item.title.rstrip(".")},"' if item.title else '"Untitled Document,"'
    year_str = f"{item.year}." if item.year else "n.d."

    doi_str = ""
    if item.doi:
        doi_str = f" doi: {item.doi}."

    return f"[{index}] {author_str}, {title_str} {year_str}{doi_str}"


class BaseMCPTransport(ABC):
    """Protocol for sending JSON-RPC messages to an MCP server."""

    @abstractmethod
    def send_request(self, method: str, params: dict[str, Any]) -> dict[str, Any]:
        """Dispatch a JSON-RPC request and return the result dictionary."""

    @abstractmethod
    def close(self) -> None:
        """Close connection to server."""


class StdioMCPTransport(BaseMCPTransport):
    """Spawns and communicates with an MCP server via stdio streams."""

    def __init__(self, command: list[str], timeout: float = 10.0) -> None:
        self.command = command
        self.timeout = timeout
        self._process: subprocess.Popen[str] | None = None
        self._request_id = 0

    def _ensure_process(self) -> subprocess.Popen[str]:
        if self._process is None or self._process.poll() is not None:
            cmd_binary = self.command[0]
            if not shutil.which(cmd_binary):
                raise MCPConnectionError(f"MCP server executable '{cmd_binary}' not found on PATH.")
            try:
                self._process = subprocess.Popen(
                    self.command,
                    stdin=subprocess.PIPE,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                    bufsize=1,
                )
            except OSError as exc:
                raise MCPConnectionError(f"Failed to launch MCP server process: {exc}") from exc
        return self._process

    def send_request(self, method: str, params: dict[str, Any]) -> dict[str, Any]:
        proc = self._ensure_process()
        self._request_id += 1
        payload = {
            "jsonrpc": "2.0",
            "id": self._request_id,
            "method": method,
            "params": params,
        }
        try:
            assert proc.stdin is not None
            proc.stdin.write(json.dumps(payload) + "\n")
            proc.stdin.flush()

            assert proc.stdout is not None
            response_line = proc.stdout.readline()
            if not response_line:
                raise MCPConnectionError("MCP server process closed stream unexpectedly.")

            data = json.loads(response_line)
            if "error" in data:
                raise MCPError(f"MCP server returned error: {data['error']}")
            return data.get("result", {})
        except (OSError, json.JSONDecodeError) as exc:
            raise MCPConnectionError(f"MCP communication failure: {exc}") from exc

    def close(self) -> None:
        if self._process is not None:
            try:
                self._process.terminate()
                self._process.wait(timeout=2.0)
            except Exception:
                self._process.kill()
            finally:
                self._process = None


class MockMCPTransport(BaseMCPTransport):
    """In-memory mock transport for unit testing MCP dispatch without external processes."""

    def __init__(self, fail_with: Exception | None = None) -> None:
        self.recorded_calls: list[dict[str, Any]] = []
        self.fail_with = fail_with

    def send_request(self, method: str, params: dict[str, Any]) -> dict[str, Any]:
        if self.fail_with:
            raise self.fail_with
        self.recorded_calls.append({"method": method, "params": params})
        tool_name = params.get("name", "")
        return {
            "content": [{"type": "text", "text": f"Successfully executed tool: {tool_name}"}],
            "status": "success",
        }

    def close(self) -> None:
        pass


class PythonDocxBackend:
    """Direct document manipulation using python-docx."""

    def __init__(self) -> None:
        self.doc: docx.Document | None = None
        self.current_path: Path | None = None
        self.references: list[CitationItem] = []

    def create_document(self, title: str, path: Path) -> Path:
        self.doc = docx.Document()
        self.current_path = path
        self.references.clear()

        heading = self.doc.add_heading(title, level=0)
        heading.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        return path

    def append_heading(self, text: str, level: int = 1) -> None:
        if self.doc is None:
            raise DocumentGenerationError("No active document. Call create_document first.")
        safe_level = max(1, min(level, 9))
        self.doc.add_heading(text, level=safe_level)

    def add_paragraph_with_citation(
        self,
        text: str,
        citations: list[CitationItem] | None = None,
        style: CitationStyle = "apa",
    ) -> None:
        if self.doc is None:
            raise DocumentGenerationError("No active document. Call create_document first.")

        p = self.doc.add_paragraph()
        p.add_run(text)

        if citations:
            if style == "apa":
                citation_text = format_inline_citation_apa(citations)
                p.add_run(f" {citation_text}")
                for c in citations:
                    if c not in self.references:
                        self.references.append(c)
            elif style == "ieee":
                indices: list[int] = []
                for c in citations:
                    if c not in self.references:
                        self.references.append(c)
                    indices.append(self.references.index(c) + 1)
                citation_text = format_inline_citation_ieee(indices)
                p.add_run(f" {citation_text}")

    def append_reference_list(
        self,
        references: list[CitationItem] | None = None,
        style: CitationStyle = "apa",
    ) -> None:
        if self.doc is None:
            raise DocumentGenerationError("No active document. Call create_document first.")

        entries = references or self.references
        if not entries:
            return

        self.doc.add_heading("References", level=1)
        for idx, item in enumerate(entries, 1):
            if style == "apa":
                ref_text = format_reference_entry_apa(item)
            else:
                ref_text = format_reference_entry_ieee(item, idx)

            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = docx.shared.Inches(0.5)
            p.paragraph_format.first_line_indent = docx.shared.Inches(-0.5)
            p.add_run(ref_text)

    def save(self, path: Path | None = None) -> Path:
        if self.doc is None:
            raise DocumentGenerationError("No active document to save.")
        target_path = path or self.current_path
        if target_path is None:
            raise DocumentGenerationError("No output path specified for document.")

        target_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            self.doc.save(str(target_path))
            return target_path
        except Exception as exc:
            raise DocumentGenerationError(f"Failed to save document to '{target_path}': {exc}") from exc


class WordCraftClient:
    """Model Context Protocol (MCP) client with python-docx fallback for academic documents."""

    def __init__(
        self,
        workspace: WorkspaceManager | None = None,
        command: list[str] | None = None,
        transport: BaseMCPTransport | None = None,
        fallback_to_docx: bool = True,
    ) -> None:
        self.workspace = workspace or WorkspaceManager()
        self.command = command or ["wordcraft-cli", "mcp"]
        self.transport = transport
        self.fallback_to_docx = fallback_to_docx
        self._docx_backend = PythonDocxBackend()
        self._mcp_active = False
        self._current_path: Path | None = None

        if self.transport is not None:
            self._mcp_active = True
        elif shutil.which(self.command[0]):
            try:
                self.transport = StdioMCPTransport(self.command)
                self._mcp_active = True
            except MCPConnectionError:
                self._mcp_active = False

    @property
    def is_using_mcp(self) -> bool:
        """Indicate whether calls are actively routed through an MCP transport."""
        return self._mcp_active and self.transport is not None

    def call_tool(self, tool_name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        """Dispatch a tool call to the MCP server or fallback to python-docx."""
        if self.is_using_mcp:
            try:
                assert self.transport is not None
                return self.transport.send_request(
                    "tools/call",
                    {"name": tool_name, "arguments": arguments},
                )
            except (MCPError, MCPConnectionError) as exc:
                if not self.fallback_to_docx:
                    raise
                # Switch to docx backend upon MCP failure
                self._mcp_active = False

        # Execute on local python-docx backend
        return self._execute_local_tool(tool_name, arguments)

    def _execute_local_tool(self, tool_name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        if tool_name == "create_document":
            path = Path(arguments["path"])
            self._docx_backend.create_document(arguments["title"], path)
            return {"status": "created", "path": str(path)}

        if tool_name == "append_heading":
            self._docx_backend.append_heading(arguments["text"], arguments.get("level", 1))
            return {"status": "heading_appended"}

        if tool_name == "add_paragraph_with_citation":
            citations = [CitationItem.from_any(c) for c in arguments.get("citations", [])]
            self._docx_backend.add_paragraph_with_citation(
                text=arguments["text"],
                citations=citations,
                style=arguments.get("style", "apa"),
            )
            return {"status": "paragraph_added"}

        if tool_name == "append_reference_list":
            refs = [CitationItem.from_any(r) for r in arguments.get("references", [])]
            self._docx_backend.append_reference_list(
                references=refs,
                style=arguments.get("style", "apa"),
            )
            return {"status": "references_appended"}

        if tool_name == "save_document":
            path = Path(arguments["path"]) if "path" in arguments else None
            saved_path = self._docx_backend.save(path)
            return {"status": "saved", "path": str(saved_path)}

        raise DocumentGenerationError(f"Unrecognized document tool: '{tool_name}'")

    def create_document(self, title: str, path: str | Path | None = None) -> Path:
        """Create a new academic document with a centered title."""
        dest_path = (
            Path(path).resolve()
            if path
            else self.workspace.vault_dir / f"{sanitize_filename(title)}.docx"
        )
        self._current_path = dest_path

        # If MCP is active, synchronize both MCP and local backend for seamless saving
        self._docx_backend.create_document(title, dest_path)
        self.call_tool("create_document", {"title": title, "path": str(dest_path)})
        return dest_path

    def append_heading(self, text: str, level: int = 1) -> None:
        """Append a section heading of specified level to document."""
        if self.is_using_mcp:
            self._docx_backend.append_heading(text, level)
        self.call_tool("append_heading", {"text": text, "level": level})

    def add_paragraph_with_citation(
        self,
        text: str,
        citations: list[Any] | None = None,
        style: CitationStyle = "apa",
    ) -> None:
        """Add a paragraph of text with optional inline citations (APA or IEEE)."""
        raw_citations = [
            CitationItem.from_any(c).model_dump()
            for c in (citations or [])
        ]
        if self.is_using_mcp:
            citation_items = [CitationItem.from_any(c) for c in (citations or [])]
            self._docx_backend.add_paragraph_with_citation(text, citation_items, style)

        self.call_tool(
            "add_paragraph_with_citation",
            {"text": text, "citations": raw_citations, "style": style},
        )

    def append_reference_list(
        self,
        references: list[Any],
        style: CitationStyle = "apa",
    ) -> None:
        """Append formatted bibliography references in APA 7th or IEEE format."""
        raw_refs = [CitationItem.from_any(r).model_dump() for r in references]
        if self.is_using_mcp:
            ref_items = [CitationItem.from_any(r) for r in references]
            self._docx_backend.append_reference_list(ref_items, style)

        self.call_tool(
            "append_reference_list",
            {"references": raw_refs, "style": style},
        )

    def save(self, path: str | Path | None = None) -> Path:
        """Save the generated .docx file to disk."""
        target_path = Path(path).resolve() if path else self._current_path
        if target_path is None:
            target_path = self.workspace.vault_dir / "document.docx"

        self.call_tool("save_document", {"path": str(target_path)})
        return self._docx_backend.save(target_path)

    def close(self) -> None:
        """Clean up MCP transport connection."""
        if self.transport:
            self.transport.close()
            self._mcp_active = False

    def __enter__(self) -> WordCraftClient:
        return self

    def __exit__(self, *args: Any) -> None:
        self.close()
